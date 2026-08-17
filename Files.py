"""Excel 文件操作（openpyxl 后端）。

对外接口是一组"扁平函数"（不需要先 new 一个类的实例），供 run_engine.py 的
8 个 excel_xxx 步骤直接调用。这一层调用约定保持不变，但函数内部的具体算法
参照了 rpaframework 项目 RPA.Excel.Files 模块（见 _internal/Files.py 的
XlsxWorkbook / Files 类）的真实实现——那是经过验证、被广泛使用的开源代码。
凡是 rpaframework 依赖 RPA.Tables / RPA.application 等本项目未安装的重型
依赖的外围功能（如 read_worksheet_as_table），本模块不涉及；但核心的
openpyxl 操作逻辑（尤其是一些容易被忽略的边界情况）尽量原样照搬。
"""
import pathlib
import re

import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.styles import Font, PatternFill, Side, Border
from openpyxl.drawing.image import Image as XLImage

try:
    from PIL import ImageColor
except ImportError:  # pragma: no cover - Pillow 是 openpyxl 图片功能的依赖，正常应已随其安装
    ImageColor = None


# ---------------- 工作簿生命周期 ----------------

def new_workbook():
    """新建一个空白工作簿（自带一个默认工作表）。"""
    return openpyxl.Workbook()


def open_workbook(path, data_only=False, read_only=False):
    """打开一个已存在的工作簿。

    - 路径不存在时会抛出清晰的 FileNotFoundError（而不是 openpyxl 内部更晦涩的报错），
      做法参照 rpaframework `_load_workbook` 里 `Path(path).resolve(strict=True)` 的用法。
    - .xlsm/.xltm（含宏的工作簿）会自动加上 keep_vba=True，否则另存后宏代码会丢失——
      这是 rpaframework `XlsxWorkbook.open()` 里专门处理的一个坑。
    - data_only=True 时，公式单元格读到的是"文件里缓存的上次计算结果"而不是公式文本本身；
      本流程工具默认 False（读公式文本），需要读计算结果请走 Application.read_calculated_value。
    """
    resolved = pathlib.Path(path).resolve(strict=True)
    suffix = resolved.suffix.lower()
    options = {"filename": str(resolved), "data_only": data_only, "read_only": read_only}
    if suffix in (".xlsm", ".xltm"):
        options["keep_vba"] = True
    return openpyxl.load_workbook(**options)


def save_workbook(wb, path):
    wb.save(str(path))


def close_workbook(wb):
    wb.close()


# ---------------- 工作表（sheet）管理 ----------------

def add_sheet(wb, name, index=None):
    """新建工作表；若同名工作表已存在则报错（参照 rpaframework `create_worksheet` 的重名检查）。"""
    name = str(name)
    if name in wb.sheetnames:
        raise ValueError(f"工作表 {name!r} 已存在")
    wb.create_sheet(name, index)


def delete_sheet(wb, name):
    """删除工作表；工作簿至少要保留一个工作表，删最后一个会报错
    （参照 rpaframework `remove_worksheet` 里 `Workbook must have at least one other worksheet` 的保护）。
    """
    if name not in wb.sheetnames:
        raise ValueError(f"工作表 {name!r} 不存在")
    others = [s for s in wb.sheetnames if s != name]
    if not others:
        raise ValueError("工作簿必须至少保留一个工作表，不能删除最后一个工作表")
    del wb[name]


def rename_sheet(wb, old_name, new_name):
    wb[old_name].title = str(new_name)


def get_sheet(wb, name):
    return wb[name] if name else wb.active


# ---------------- 整表读写 ----------------

def read_all(ws):
    return [list(row) for row in ws.iter_rows(values_only=True)]


def write_all(ws, data, start_cell="A1"):
    start_row, start_col = coordinate_to_tuple(start_cell)
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            ws.cell(row=start_row + r, column=start_col + c, value=val)


def clear_all(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None


# ---------------- 单元格 / 区域读写 ----------------

def read_cell(ws, ref):
    return ws[ref].value


def write_cell(ws, ref, value):
    ws[ref] = value


def read_range(ws, ref):
    return [[cell.value for cell in row] for row in ws[ref]]


def write_range(ws, start_ref, data):
    start_row, start_col = coordinate_to_tuple(start_ref)
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            ws.cell(row=start_row + r, column=start_col + c, value=val)


# ---------------- 行列结构编辑 ----------------

def insert_rows(ws, row_idx, count=1):
    ws.insert_rows(row_idx, count)


def delete_rows(ws, row_idx, count=1):
    ws.delete_rows(row_idx, count)


def _col_index(col_ref):
    """列既可能是字母（'C'）也可能已经是数字下标，统一转成数字下标。"""
    return col_ref if isinstance(col_ref, int) else column_index_from_string(str(col_ref))


def insert_cols(ws, col_ref, count=1):
    ws.insert_cols(_col_index(col_ref), count)


def delete_cols(ws, col_ref, count=1):
    ws.delete_cols(_col_index(col_ref), count)


# ---------------- 样式格式 ----------------

def _iter_cells(ws, ref):
    for row in ws[ref]:
        for cell in (row if hasattr(row, '__iter__') else [row]):
            yield cell


_HEX_COLOR_RE = re.compile(r"^(?:[0-9a-fA-F]{3}){2}$")


def _to_hex_color(value):
    """颜色既可以传十六进制（'FF0000'）也可以传颜色名（'red'），
    参照 rpaframework `_set_color_if_given`/`_set_fill_color` 里先判断是否已是十六进制、
    否则用 PIL.ImageColor 解析颜色名再转十六进制的做法。
    """
    if not value:
        return None
    value = str(value).strip()
    if _HEX_COLOR_RE.match(value):
        return value
    if ImageColor is not None:
        try:
            r, g, b = ImageColor.getrgb(value)
            return "%02x%02x%02x" % (r, g, b)
        except ValueError:
            pass
    return value


def set_font(ws, ref, name=None, size=None, bold=None, color=None):
    color_hex = _to_hex_color(color)
    for cell in _iter_cells(ws, ref):
        kw = {}
        if name is not None:
            kw['name'] = name
        if size is not None and size > 0:
            kw['size'] = size
        if bold is not None:
            kw['bold'] = bold
        if color_hex:
            kw['color'] = color_hex
        cell.font = Font(**kw)


def set_fill_color(ws, ref, color_hex):
    color_hex = _to_hex_color(color_hex)
    fill = PatternFill(fill_type='solid', fgColor=color_hex)
    for cell in _iter_cells(ws, ref):
        cell.fill = fill


def set_border(ws, ref, style="thin"):
    side = Side(border_style=None) if style == 'none' else Side(border_style=style)
    border = Border(left=side, right=side, top=side, bottom=side)
    for cell in _iter_cells(ws, ref):
        cell.border = border


# ---------------- 公式 ----------------

def write_formula(ws, ref, formula):
    ws[ref] = formula


# ---------------- 图片 ----------------

def insert_image(ws, ref, image_path, width=0, height=0):
    img = XLImage(image_path)
    if width > 0:
        img.width = width
    if height > 0:
        img.height = height
    ws.add_image(img, ref)


def clear_images(ws):
    ws._images.clear()
